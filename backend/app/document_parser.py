from __future__ import annotations

import io
import re

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook


class DocumentReadError(Exception):
    pass


def read_attachment(data: bytes, filename: str) -> str:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    try:
        if ext == "txt":
            return data.decode("utf-8", errors="replace")

        if ext == "pdf":
            reader = PdfReader(io.BytesIO(data))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            if not text.strip():
                raise DocumentReadError("PDF contained no extractable text")
            return text

        if ext == "docx":
            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    vals = [cell.text.strip() for cell in row.cells]
                    if any(vals):
                        parts.append(" | ".join(vals))
            text = "\n".join(parts)
            if not text.strip():
                raise DocumentReadError("DOCX contained no extractable text")
            return text

        if ext == "xlsx":
            wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
            parts = []
            for ws in wb.worksheets:
                parts.append(f"[SHEET {ws.title}]")
                for row in ws.iter_rows(values_only=True):
                    vals = ["" if v is None else str(v) for v in row]
                    if any(vals):
                        parts.append(" | ".join(vals))
            text = "\n".join(parts)
            if not text.strip():
                raise DocumentReadError("XLSX contained no extractable text")
            return text

        raise DocumentReadError(f"Unsupported attachment type: .{ext}")
    except DocumentReadError:
        raise
    except Exception as exc:
        raise DocumentReadError(f"{type(exc).__name__}: {exc}") from exc


def detect_doc_type(text: str) -> str:
    if not text or not text.strip():
        return "UNREADABLE"
    if text.startswith("[ERROR"):
        return "UNREADABLE"

    low = text.lower()

    if "commercial invoice" in low or "invoice no." in low:
        return "INVOICE"
    if "packing list" in low:
        return "PACKING_LIST"
    if "certificate of origin" in low:
        return "CERTIFICATE"
    if (
        "shipping instruction" in low
        or "bl instruction" in low
        or "[sheet s.i.]" in low
        or "bill of lading instruction" in low
    ):
        return "SI"
    if (
        "bill of lading" in low
        or re.search(r"\bb/l\s*(?:no|number)", low)
        or re.search(r"\bbl\s+no\.", low)
    ):
        return "BL"
    return "OTHER"